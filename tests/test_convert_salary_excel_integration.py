"""Exercise the documented Excel -> converter CLI -> salary JSON path."""
import json
import shutil
import subprocess
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

try:
    import openpyxl
except ImportError:
    openpyxl = None

from salary_lookup import collect_validation_issues, format_entry

TOOL = Path(__file__).resolve().parent.parent / "tools" / "convert_salary_excel.py"


@unittest.skipUnless(openpyxl, "requires optional openpyxl (installed in CI)")
class ConverterCLI(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.workbook = self.root / "salary input.xlsx"
        self.output = self.root / "salary output.json"

    def run_converter(self, *args):
        return subprocess.run(
            [sys.executable, str(TOOL), str(self.workbook),
             "--output", str(self.output), *args],
            cwd=self.root, capture_output=True, text=True, encoding="utf-8",
        )

    def test_real_workbook_round_trips_multiple_sheets_and_metadata(self):
        book = openpyxl.Workbook()
        try:
            first = book.active
            first.title = "Personal ficticio"
            first.append(["Empresa", "Ciudad", "Cantidad", "Índice"])
            first.append(["Ñandú Ficticia SRL", "Córdoba", 500, 108.5])
            first.append(["Compañía Demo SA", "Neuquén", 12, "102,5"])
            second = book.create_sheet("Desarrollo ficticio")
            second.append(["Empresa", "Ciudad", "Desarrollo Cantidad", "Desarrollo Índice"])
            second.append(["Empresa Prueba SAS", "Rosario", 8, 112.5])
            book.save(self.workbook)
        finally:
            book.close()

        result = self.run_converter(
            "--source", "Integration fixture", "--baseline", "100",
            "--baseline-desc", "Index 100 = median salary",
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads(self.output.read_text(encoding="utf-8"))
        self.assertEqual(data["metadata"], {
            "source": "Integration fixture", "index_baseline": 100,
            "index_label": "Index", "baseline_description": "Index 100 = median salary",
        })
        self.assertEqual(data["companies"], [
            {"company": "Ñandú Ficticia SRL", "city": "Córdoba", "categories": {
                "all_employees": {"count": 500, "index": 108.5}}},
            {"company": "Compañía Demo SA", "city": "Neuquén", "categories": {
                "all_employees": {"count": 12, "index": 102.5}}},
            {"company": "Empresa Prueba SAS", "city": "Rosario", "categories": {
                "desarrollo": {"count": 8, "index": 112.5}}},
        ])
        self.assertEqual(collect_validation_issues(data), ([], []))
        rendered = format_entry(data["companies"][0], data["metadata"])
        self.assertIn("All Employees", rendered)
        self.assertIn("+8.5%", rendered)
        self.assertNotIn("N/A*", rendered)

    def test_workbook_without_salary_headers_fails_without_writing_output(self):
        book = openpyxl.Workbook()
        try:
            book.active.append(["Notes", "Comments"])
            book.active.append(["No salary data", "Just a worksheet"])
            book.save(self.workbook)
        finally:
            book.close()

        result = self.run_converter()
        self.assertEqual(result.returncode, 1, result.stderr)
        self.assertIn("No data could be parsed", result.stderr)
        self.assertFalse(self.output.exists())

    def test_spanish_amount_workbook_to_json_lookup_keeps_explicit_units(self):
        book = openpyxl.Workbook()
        try:
            book.active.append(["Empresa", "Ciudad", "Sueldo", "CUIT", "Aumento %"])
            book.active.append(["Ñandú Ficticia S.R.L.", "Córdoba", "123.456,78", 12345678901, 15])
            book.save(self.workbook)
        finally:
            book.close()
        result = self.run_converter("--currency", "ARS", "--basis", "bruto", "--period", "mensual", "--source", "Demostración ficticia", "--as-of", "2026-10-01")
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads(self.output.read_text(encoding="utf-8"))
        category = data["companies"][0]["categories"]["sueldo"]
        self.assertEqual(category["amount"], 123456.78)
        self.assertEqual(category["currency"], "ARS")
        self.assertEqual(category["pay_basis"], "gross")
        self.assertEqual(category["period"], "month")
        self.assertEqual(set(data["companies"][0]["categories"]), {"sueldo"})
        self.assertNotIn("index", category)
        self.assertEqual(collect_validation_issues(data), ([], []))
        shutil.copy(TOOL.parent.parent / "salary_lookup.py", self.root / "salary_lookup.py")
        shutil.copy(self.output, self.root / "salary_data.json")
        lookup = subprocess.run([sys.executable, str(self.root / "salary_lookup.py"), "Nandu Ficticia", "--city", "Cordoba", "--json"], capture_output=True, text=True, encoding="utf-8", cwd=self.root)
        self.assertEqual(lookup.returncode, 0, lookup.stderr)
        self.assertEqual(json.loads(lookup.stdout)[0]["categories"]["sueldo"]["as_of"], "2026-10-01")

    def test_missing_amount_freshness_fails_without_replacing_previous_output(self):
        book = openpyxl.Workbook()
        try:
            book.active.append(["Empresa", "Salario"])
            book.active.append(["Empresa Ficticia SA", 123])
            book.save(self.workbook)
        finally:
            book.close()
        self.output.write_bytes(b"previous user-managed data")
        result = self.run_converter("--currency", "USD", "--basis", "neto", "--period", "anual", "--source", "Demo")
        self.assertEqual(result.returncode, 1)
        self.assertIn("as_of", result.stderr)
        self.assertEqual(self.output.read_bytes(), b"previous user-managed data")

    def test_mixed_currencies_and_periods_from_explicit_excel_columns(self):
        book = openpyxl.Workbook()
        try:
            book.active.append(["Razón social", "Ciudad", "Salario", "Moneda", "Bruto/Neto", "Período", "Fuente", "Fecha"])
            book.active.append(["Primera Ficticia SA", "Córdoba", 123, "ARS", "bruto", "mensual", "Fixture", "2026-10-01"])
            book.active.append(["Segunda Ficticia SAS", "Neuquén", 456, "USD", "neto", "anual", "Fixture", "2026-10-01"])
            book.save(self.workbook)
        finally:
            book.close()
        result = self.run_converter()
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads(self.output.read_text(encoding="utf-8"))
        local, foreign = [item["categories"]["salario"] for item in data["companies"]]
        self.assertEqual((local["amount"], local["currency"], local["period"]), (123, "ARS", "month"))
        self.assertEqual((foreign["amount"], foreign["currency"], foreign["period"]), (456, "USD", "year"))


if __name__ == "__main__":
    unittest.main()
