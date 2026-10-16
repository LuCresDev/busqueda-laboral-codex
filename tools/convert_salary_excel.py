#!/usr/bin/env python3
"""
Convertir datos salariales propios de Excel a JSON.

This script converts an Excel file containing company salary data
into the JSON format expected by salary_lookup.py.

Prerequisites:
    pip install openpyxl

Usage:
    python tools/convert_salary_excel.py <path-to-excel-file>
    python tools/convert_salary_excel.py <path-to-excel-file> --source "Fuente propia"
    python tools/convert_salary_excel.py <path-to-excel-file> --baseline 100 --baseline-desc "Index 100 = median salary"

The output file (salary_data.json) will be written to the repository root.

Expected Excel format:
    - A header row with column names
    - Una columna Empresa/Compañía/Razón social (required)
    - Una columna Ciudad/Localidad (optional)
    - Any number of numeric data columns (salary index, count, etc.)

The script auto-detects the header row and column layout. For Excel files
with paired count/index columns per category, it groups them automatically.
"""

import json
import sys
import argparse
import re
import math
import unicodedata
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    openpyxl = None


# Column name patterns for auto-detection
COMPANY_PATTERNS = {"empresa", "compañía", "compania", "razón", "razon", "firma", "company", "virksomhed", "employer", "arbejdsgiver"}
CITY_PATTERNS = {"ciudad", "localidad", "ubicación", "ubicacion", "by", "city", "kommune", "location", "lokation", "sted"}
COUNT_PATTERNS = {"cantidad", "empleados", "personas", "muestra", "antal", "count", "number", "n", "employees", "medarbejdere"}
STRICT_INDEX_PATTERNS = {"índice", "indice", "indeks", "index", "idx"}
INDEX_PATTERNS = STRICT_INDEX_PATTERNS | {"salary", "løn", "median", "average", "gennemsnit"}
AMOUNT_PATTERNS = {"sueldo", "sueldos", "salario", "salarios", "remuneración", "remuneracion", "haberes"}
PERCENT_PATTERNS = {"porcentaje", "porcentual", "percent", "percentage", "aumento", "variación", "variacion"}
METADATA_PATTERNS = {
    "currency": {"moneda", "currency"},
    "pay_basis": {"bruto", "neto", "basis"},
    "period": {"período", "periodo", "period", "frecuencia"},
    "source": {"fuente", "source"},
    "as_of": {"fecha", "date"},
}
# "Compound" tokens: pattern words allowed to match as a substring of a larger
# header token, for languages that glue words together (e.g. Danish "lønindeks"
# -> løn + indeks). Languages that write headers as separate words need none.
# Conserva compatibilidad con encabezados compuestos de fuentes históricas
# multilingües; no establece un país, una moneda ni una fuente predeterminados.
COMPOUND_PATTERNS = {"antal", "indeks", "løn", "gennemsnit", "medarbejdere"}
# Identifier columns (employee id, Danish "personnummer", etc.) are never salary
# data. They are dropped at classification so they are not mistaken for a salary
# category. Matched as whole tokens only, like other pattern sets.
ID_PATTERNS = {"id", "personnummer", "dni", "cuit", "cuil", "legajo", "identificador", "código", "codigo"}
# Category name for a count/index pair whose headers carry no category word at
# all ("Count" + "Index"). Matches the top-level category in README_SALARY_TOOL.md.
DEFAULT_CATEGORY = "all_employees"


def parse_numeric_cell(value):
    """Parse numeric Excel values, including localized string cells."""
    if isinstance(value, bool):
        raise ValueError("boolean is not salary data")
    if isinstance(value, (int, float)):
        number = float(value)
        if not math.isfinite(number):
            raise ValueError("not finite")
        return number
    if not isinstance(value, str):
        raise ValueError("not numeric")

    text = value.strip().replace("\u00a0", " ").replace(" ", "")
    if not text:
        raise ValueError("not numeric")
    if "," in text and "." in text:
        # The separator that appears last is the decimal separator: European
        # "1.234,56" and US "1,234.56" are both unambiguous here, unlike the
        # single-separator cases below.
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        if re.fullmatch(r"[+-]?\d+,\d{3}", text):
            raise ValueError("ambiguous comma separator")
        text = text.replace(",", ".")
    elif "." in text:
        if re.fullmatch(r"[+-]?\d+\.\d{3}", text):
            raise ValueError("ambiguous dot separator")
    number = float(text)
    if not math.isfinite(number):
        raise ValueError("not finite")
    return number


def parse_count_cell(value):
    try:
        number = parse_numeric_cell(value)
    except (TypeError, ValueError):
        return None
    if number < 0 or not number.is_integer():
        raise ValueError("Count must be a non-negative integer; no fractional count is truncated")
    return int(number)


def header_matches(header, patterns):
    """Return True when a header contains a meaningful pattern match.

    Patterns match whole tokens; any pattern also listed in
    ``COMPOUND_PATTERNS`` may additionally match as a substring, to handle
    languages that form compound words.
    """
    h = unicodedata.normalize("NFC", header).lower().strip()
    tokens = set(re.findall(r"[^\W_]+", h))

    for p in patterns:
        if p in tokens:
            return True
        if p in COMPOUND_PATTERNS and p in h:
            return True
    return False


def strip_type_patterns(header, patterns):
    """Remove count/index words from a header to derive a category name.

    Mirrors ``header_matches``: patterns strip as whole tokens, and any
    pattern also listed in ``COMPOUND_PATTERNS`` additionally strips as a
    substring - otherwise a compound header like "Lønindeks alle" keeps the
    type word in its category name and can never pair with "Antal alle".
    """
    name = unicodedata.normalize("NFC", header).lower()
    for p in patterns:
        name = re.sub(rf"(?<![^\W_]){re.escape(p)}(?![^\W_])", "", name)
        if p in COMPOUND_PATTERNS:
            name = name.replace(p, "")
    return name.strip(" _-")


def detect_column_type(header):
    """Detect whether a column header refers to count or index data."""
    if header_matches(header, ID_PATTERNS) or "%" in header or header_matches(header, PERCENT_PATTERNS):
        return None
    if header_matches(header, COUNT_PATTERNS):
        return "count"
    if header_matches(header, STRICT_INDEX_PATTERNS):
        return "index"
    if header_matches(header, AMOUNT_PATTERNS):
        return "amount"
    if header_matches(header, INDEX_PATTERNS):
        return "index"
    return None


def _context_value(key, value):
    if key == "as_of" and isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    value = str(value).strip()
    if key == "currency":
        return value.upper()
    if key == "pay_basis":
        return {"bruto": "gross", "neto": "net"}.get(value.lower(), value.lower())
    if key == "period":
        return {"mensual": "month", "anual": "year", "hora": "hour", "semanal": "week", "diario": "day"}.get(value.lower(), value.lower())
    return value


def parse_sheet(ws, sheet_label=None, salary_metadata=None, value_kind="legacy"):
    """Parse a single worksheet into a list of company entries and detected categories."""
    # Find header row. Two passes:
    #
    # Strict pass: a candidate row needs a company-pattern cell AND a
    # DIFFERENT cell matching a city/count/index pattern. Corroboration must
    # come from a separate cell - a single free-text sentence can pack both
    # a company-pattern word and a count-pattern word together (e.g. "...
    # opdelt efter arbejdsgiver, antal svar 1234"), and that must not read
    # as a header any more than a citation mentioning just one of them does.
    # A real header row always has these as separate columns.
    #
    # Fallback pass: some real headers have no recognizable city/count/index
    # column at all (e.g. "Company | Base pay 2025 | Bonus 2025" - neither
    # data header matches a known pattern, so they're picked up later as
    # untyped/standalone categories). Nothing can corroborate a company match
    # there, so if the strict pass finds no row in the first 10, fall back to
    # the original any-cell-mentions-company rule.
    rows = list(ws.iter_rows(min_row=1, max_row=10, values_only=False))

    def _cell_texts(row):
        return [str(cell.value).strip() for cell in row if cell.value]

    header_row = None
    for row_idx, row in enumerate(rows, start=1):
        cell_texts = _cell_texts(row)
        company_idxs = {i for i, t in enumerate(cell_texts) if header_matches(t, COMPANY_PATTERNS) and not header_matches(t, ID_PATTERNS)}
        if not company_idxs:
            continue
        other_idxs = {
            i
            for i, t in enumerate(cell_texts)
            if header_matches(t, CITY_PATTERNS) or detect_column_type(t) is not None
        }
        if other_idxs - company_idxs:
            header_row = row_idx
            break

    if header_row is None:
        for row_idx, row in enumerate(rows, start=1):
            if any(header_matches(t, COMPANY_PATTERNS) for t in _cell_texts(row)):
                header_row = row_idx
                break

    if header_row is None:
        print(f"Warning: Could not find header row in sheet '{ws.title}'. Skipping.", file=sys.stderr)
        return []

    # Read headers
    headers = []
    for cell in ws[header_row]:
        headers.append(str(cell.value).strip() if cell.value else "")

    # Find company and city columns
    company_col = None
    city_col = None
    for i, h in enumerate(headers):
        if header_matches(h, ID_PATTERNS):
            continue
        if header_matches(h, COMPANY_PATTERNS):
            company_col = i
        elif header_matches(h, CITY_PATTERNS):
            city_col = i

    if company_col is None:
        print(f"Warning: Could not find company column in sheet '{ws.title}'.", file=sys.stderr)
        return []

    # Identify data columns (everything that's not company/city or an identifier)
    data_cols = []
    metadata_cols = {}
    for i, h in enumerate(headers):
        if i == company_col or i == city_col or not h:
            continue
        if header_matches(h, ID_PATTERNS) or "%" in h or header_matches(h, PERCENT_PATTERNS):
            continue
        context_key = next((key for key, patterns in METADATA_PATTERNS.items() if header_matches(h, patterns) and not header_matches(h, AMOUNT_PATTERNS | INDEX_PATTERNS)), None)
        if context_key:
            metadata_cols[context_key] = i
            continue
        data_cols.append((i, h))

    # Group data columns by detected type and derive category names
    count_cols = []
    index_cols = []
    untyped_cols = []
    amount_cols = []

    for col_idx, col_header in data_cols:
        col_type = detect_column_type(col_header)
        if col_type == "index" and value_kind == "amount" and not header_matches(col_header, STRICT_INDEX_PATTERNS):
            col_type = "amount"
        if col_type == "count":
            cat_name = strip_type_patterns(col_header, COUNT_PATTERNS)
            count_cols.append((col_idx, col_header, cat_name))
        elif col_type == "index":
            cat_name = strip_type_patterns(col_header, INDEX_PATTERNS)
            index_cols.append((col_idx, col_header, cat_name))
        elif col_type == "amount":
            amount_cols.append((col_idx, col_header))
        else:
            untyped_cols.append((col_idx, col_header))

    # Pair count/index columns by matching category name. A bare "Count" /
    # "Index" pair (Danish "Antal" / "Lønindeks") strips to an empty name on
    # both sides - the single-category layout the README's "auto-pairs
    # count/index columns" line describes. It is still one pair, so it gets
    # the README's default category name instead of being emitted as two
    # unrelated standalone columns: salary_lookup renders that split as a
    # count row whose index reads "N/A*", i.e. "too few employees to publish
    # (privacy)", about a company with a published headcount.
    categories = []
    used_counts = set()
    used_indexes = set()

    for ci, (c_idx, c_header, c_cat) in enumerate(count_cols):
        for ii, (i_idx, i_header, i_cat) in enumerate(index_cols):
            if ii in used_indexes:
                continue
            if c_cat == i_cat:
                cat_name = (c_cat or DEFAULT_CATEGORY).replace(" ", "_").replace("-", "_")
                categories.append({
                    "name": cat_name,
                    "count_col": c_idx,
                    "index_col": i_idx,
                })
                used_counts.add(ci)
                used_indexes.add(ii)
                break

    # Remaining unmatched count columns become standalone. They are still count
    # data, so tag them as such — otherwise a lone headcount would be emitted as
    # a salary index and rendered with a meaningless "vs baseline" percentage.
    for ci, (c_idx, c_header, _) in enumerate(count_cols):
        if ci not in used_counts:
            categories.append(
                {"name": c_header.lower().replace(" ", "_"), "value_col": c_idx, "field": "count"}
            )

    # Remaining unmatched index columns become standalone (use original header)
    for ii, (i_idx, i_header, _) in enumerate(index_cols):
        if ii not in used_indexes:
            categories.append({"name": i_header.lower().replace(" ", "_"), "value_col": i_idx})

    # Untyped columns become standalone
    for col_idx, col_header in untyped_cols:
        categories.append({"name": col_header.lower().replace(" ", "_"), "value_col": col_idx})
    for col_idx, col_header in amount_cols:
        categories.append({"name": col_header.lower().replace(" ", "_"), "value_col": col_idx, "field": "amount"})

    if not categories:
        print(
            f"Warning: No salary data columns detected in sheet '{ws.title}' "
            "(only a company/city column was found) - the header row may be "
            "wrong, or this sheet has no salary data.",
            file=sys.stderr,
        )

    # Parse data rows
    companies = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if company_col >= len(row) or not row[company_col]:
            continue

        company_name = str(row[company_col]).strip()
        if city_col is not None and city_col < len(row) and row[city_col]:
            city_name = str(row[city_col]).strip()
        else:
            city_name = ""

        entry = {
            "company": company_name,
            "city": city_name,
            "categories": {},
        }
        context = dict(salary_metadata or {})
        for key, col_idx in metadata_cols.items():
            if col_idx < len(row) and row[col_idx] is not None:
                supplied = _context_value(key, row[col_idx])
                if key in context and context[key] is not None and context[key] != supplied:
                    raise ValueError(f"Conflicting {key} in sheet {ws.title}: row metadata differs from command metadata")
                context[key] = supplied

        for cat in categories:
            cat_name = cat["name"]
            if "count_col" in cat and "index_col" in cat:
                count_val = None
                index_val = None
                if cat["count_col"] < len(row) and row[cat["count_col"]] is not None:
                    count_val = parse_count_cell(row[cat["count_col"]])
                if cat["index_col"] < len(row) and row[cat["index_col"]] is not None:
                    try:
                        index_val = parse_numeric_cell(row[cat["index_col"]])
                    except (ValueError, TypeError):
                        pass
                # A count/index pair that is entirely empty for this row carries
                # no salary information, so skip it rather than emit nulls.
                if count_val is None and index_val is None:
                    continue
                entry["categories"][cat_name] = {"count": count_val, "index": index_val}
            elif "value_col" in cat:
                if cat["value_col"] < len(row) and row[cat["value_col"]] is not None:
                    val = row[cat["value_col"]]
                    try:
                        val = parse_numeric_cell(val)
                    except (ValueError, TypeError):
                        # Non-numeric standalone value (e.g. a free-text "Notes"
                        # column) is not salary data; skip it for this row.
                        continue
                    field = cat.get("field", "index")
                    if field == "count" and (val < 0 or not val.is_integer()):
                        raise ValueError(f"Count must be a non-negative integer in sheet {ws.title}: {cat_name}")
                    entry["categories"][cat_name] = {field: int(val) if field == "count" else val}
                    if field == "amount":
                        entry["categories"][cat_name].update({key: value for key, value in context.items() if key in METADATA_PATTERNS and value is not None})

        companies.append(entry)

    return companies


def _force_utf8_output() -> None:
    """Write UTF-8 whatever the host's default encoding is.

    A piped stdout on Windows defaults to the ANSI code page (cp1252 on most
    Western installs), so printing a company, title or file name outside it
    raised UnicodeEncodeError before the workflow saw any output.
    """
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)  # absent on a StringIO under test
        if reconfigure:
            reconfigure(encoding="utf-8")


def main():
    _force_utf8_output()
    parser = argparse.ArgumentParser(
        description="Convert salary Excel data to JSON"
    )
    parser.add_argument("excel_file", help="Path to the Excel file with salary data")
    parser.add_argument(
        "--output", default=None,
        help="Output JSON file path (default: salary_data.json in repo root)",
    )
    parser.add_argument(
        "--source", default=None,
        help="Name of the data source (e.g., 'Union Statistics 2025')",
    )
    parser.add_argument(
        "--baseline", type=float, default=100,
        help="Baseline value for index comparison (default: 100)",
    )
    parser.add_argument(
        "--baseline-desc", default=None,
        help="Description of what the baseline means (e.g., 'Index 100 = median salary')",
    )
    parser.add_argument("--currency", help="Moneda explícita para importes, por ejemplo ARS o USD; no convierte valores")
    parser.add_argument("--pay-basis", "--basis", dest="pay_basis", choices=("gross", "net", "bruto", "neto"), help="Bruto o neto para importes absolutos")
    parser.add_argument("--period", choices=("month", "year", "hour", "week", "day", "mensual", "anual", "hora", "semanal", "diario"))
    parser.add_argument("--as-of", help="Fecha del dato salarial YYYY-MM-DD, sin actualización automática")
    parser.add_argument("--value-kind", choices=("legacy", "amount"), default="legacy", help="amount interpreta encabezados salariales ingleses como importes; no cambia columnas explícitas de índice")
    args = parser.parse_args()

    excel_path = Path(args.excel_file)
    if not excel_path.exists():
        print(f"Error: File not found: {excel_path}", file=sys.stderr)
        sys.exit(1)

    if openpyxl is None:
        print("Error: openpyxl is required. Install it with: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output) if args.output else Path(__file__).parent.parent / "salary_data.json"

    print(f"Reading: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)

    salary_metadata = {key: _context_value(key, getattr(args, key)) for key in METADATA_PATTERNS if getattr(args, key, None) is not None}
    all_companies = []
    try:
        for sheet_name in wb.sheetnames:
            print(f"  Parsing sheet: {sheet_name}")
            ws = wb[sheet_name]
            companies = parse_sheet(ws, sheet_label=sheet_name, salary_metadata=salary_metadata, value_kind=args.value_kind)
            all_companies.extend(companies)
    except ValueError as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    finally:
        wb.close()

    if not all_companies:
        print("Error: No data could be parsed from the Excel file.", file=sys.stderr)
        print("Make sure the Excel file has a header row with a 'Company'/'Firma' column.", file=sys.stderr)
        sys.exit(1)

    # Build output
    output = {
        "metadata": {
            "source": args.source or excel_path.stem,
            "index_baseline": args.baseline,
            "index_label": "Index",
            "baseline_description": args.baseline_desc or f"Index {args.baseline} = baseline",
        },
        "companies": all_companies,
    }
    output["metadata"].update(salary_metadata)
    # Reuse the lookup's contract before touching an existing output file.
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from salary_lookup import collect_validation_issues
    errors, _ = collect_validation_issues(output)
    if errors:
        print("Error: " + "\n".join(errors), file=sys.stderr)
        return 1

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"\nDone! Wrote {len(all_companies)} company entries to {output_path}")


if __name__ == "__main__":
    sys.exit(main())
